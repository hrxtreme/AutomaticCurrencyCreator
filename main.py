import sys
import time
import pandas as pd
import traceback
import json
import os
import logging
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QComboBox, QLabel, QLineEdit, QPushButton, QTextEdit, QFileDialog,
    QMessageBox, QRadioButton, QGroupBox, QProgressBar, QDialog, QGridLayout
)
from PyQt6.QtGui import QFont, QIntValidator, QIcon, QAction, QDesktopServices, QPixmap
from PyQt6.QtCore import Qt, QThread, QObject, pyqtSignal, QUrl, QPoint, QSize

from playwright.sync_api import sync_playwright, Page, expect
from openpyxl import load_workbook

# Conditional import for Windows theme detection
try:
    import winreg
except ImportError:
    winreg = None

CONFIG_FILE = "config.json"
LOG_FILE = "automation_log.txt"

# =============================================================================
# THEME STYLESHEETS (QSS)
# =============================================================================

LIGHT_THEME = """
    #MainWindow, QDialog {
        background-color: #f0f2f5;
    }
    #SplashScreen #SplashFrame {
        background-color: #ffffff;
        border: 2px solid #c0c0c0;
    }
    #SplashScreen QLabel {
        color: #111;
    }
    #CustomTitleBar {
        background-color: #e4e6eb;
    }
    #TitleBarButton {
        background-color: transparent;
        border: none;
        color: #333;
        font-size: 10pt;
    }
    #TitleBarButton:hover {
        background-color: #dcdcdc;
    }
    #CloseButton:hover {
        background-color: #e81123;
        color: white;
    }
    QGroupBox {
        font-family: 'Segoe UI', sans-serif;
        font-size: 11pt;
        font-weight: bold;
        color: #111;
        border: 1px solid #dcdcdc;
        border-radius: 8px;
        margin-top: 15px;
        background-color: #ffffff;
        padding-top: 20px;
    }
    QGroupBox::title {
        subcontrol-origin: margin;
        subcontrol-position: top left;
        padding: 5px 12px;
        background-color: #f0f2f5;
        border-radius: 4px;
        border: 1px solid #dcdcdc;
        color: #333;
        margin-left: 10px;
    }
    QLabel, QRadioButton {
        font-family: 'Segoe UI', sans-serif;
        font-size: 10pt;
        color: #111;
    }
    QLineEdit, QComboBox {
        font-family: 'Segoe UI', sans-serif;
        font-size: 10pt;
        padding: 8px;
        border: 1px solid #ccc;
        border-radius: 4px;
        background-color: #fdfdfd;
        color: #111;
    }
    QLineEdit:focus, QComboBox:focus {
        border: 1px solid #0078d7;
    }
    QPushButton {
        font-family: 'Segoe UI', sans-serif;
        font-size: 10pt;
        font-weight: bold;
        padding: 8px;
        border-radius: 4px;
        border: 1px solid #ccc;
        background-color: #f5f5f5;
        color: #333;
    }
    QPushButton:hover {
        background-color: #e9e9e9;
        border-color: #bbb;
    }
    #StartButton {
        font-size: 12pt;
        padding: 12px;
        color: white;
        background-color: #27ae60;
        border: none;
    }
    #StartButton:hover {
        background-color: #2ecc71;
    }
    QTextEdit {
        font-family: 'Consolas', 'Courier New', monospace;
        font-size: 9pt;
        border: 1px solid #dcdcdc;
        border-radius: 8px;
        background-color: #ffffff;
        color: #111;
    }
"""

DARK_THEME = """
    #MainWindow, QDialog {
        background-color: #202020;
    }
    #SplashScreen #SplashFrame {
        background-color: #2d2d2d;
        border: 2px solid #444;
    }
    #SplashScreen QLabel {
        color: #e0e0e0;
    }
    #CustomTitleBar {
        background-color: #1a1a1a;
    }
    #TitleBarButton {
        background-color: transparent;
        border: none;
        color: #e0e0e0;
        font-size: 10pt;
    }
    #TitleBarButton:hover {
        background-color: #3d3d3d;
    }
    #CloseButton:hover {
        background-color: #e81123;
        color: white;
    }
    QGroupBox {
        font-family: 'Segoe UI', sans-serif;
        font-size: 11pt;
        font-weight: bold;
        color: #e0e0e0;
        border: 1px solid #444;
        border-radius: 8px;
        margin-top: 15px;
        background-color: #2d2d2d;
        padding-top: 20px;
    }
    QGroupBox::title {
        subcontrol-origin: margin;
        subcontrol-position: top left;
        padding: 5px 12px;
        background-color: #202020;
        border-radius: 4px;
        border: 1px solid #444;
        color: #bbb;
        margin-left: 10px;
    }
    QLabel, QRadioButton {
        font-family: 'Segoe UI', sans-serif;
        font-size: 10pt;
        color: #e0e0e0;
    }
    QLineEdit, QComboBox {
        font-family: 'Segoe UI', sans-serif;
        font-size: 10pt;
        padding: 8px;
        border: 1px solid #555;
        border-radius: 4px;
        background-color: #3d3d3d;
        color: #e0e0e0;
    }
    QLineEdit:focus, QComboBox:focus {
        border: 1px solid #0078d7;
    }
    QPushButton {
        font-family: 'Segoe UI', sans-serif;
        font-size: 10pt;
        font-weight: bold;
        padding: 8px;
        border-radius: 4px;
        border: 1px solid #555;
        background-color: #4a4a4a;
        color: #e0e0e0;
    }
    QPushButton:hover {
        background-color: #5a5a5a;
        border-color: #666;
    }
    #StartButton {
        font-size: 12pt;
        padding: 12px;
        color: white;
        background-color: #27ae60;
        border: none;
    }
    #StartButton:hover {
        background-color: #2ecc71;
    }
    QTextEdit {
        font-family: 'Consolas', 'Courier New', monospace;
        font-size: 9pt;
        border: 1px solid #444;
        border-radius: 8px;
        background-color: #2d2d2d;
        color: #e0e0e0;
    }
"""
# =============================================================================
# Logging Setup
# =============================================================================
def setup_logging():
    """Sets up logging to a file."""
    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s - %(message)s',
        filename=LOG_FILE,
        filemode='w'
    )
setup_logging()

# =============================================================================
# Theme Detection
# =============================================================================
def is_windows_dark_theme():
    """Checks the Windows Registry to see if the dark theme is enabled."""
    if winreg is None:
        return False
    try:
        key = winreg.OpenKey(winreg.HKEY_CURRENT_USER, r'Software\Microsoft\Windows\CurrentVersion\Themes\Personalize')
        value, _ = winreg.QueryValueEx(key, 'AppsUseLightTheme')
        return value == 0
    except (FileNotFoundError, OSError):
        return False

# =============================================================================
# Splash Screen Class
# =============================================================================
class SplashScreen(QWidget):
    def __init__(self):
        super().__init__()
        self.setObjectName("SplashScreen")
        self.setFixedSize(400, 200)
        self.setWindowFlags(Qt.WindowType.FramelessWindowHint | Qt.WindowType.WindowStaysOnTopHint)
        self.setAttribute(Qt.WidgetAttribute.WA_TranslucentBackground)
        
        layout = QVBoxLayout()
        self.setLayout(layout)

        self.frame = QWidget(self)
        self.frame.setObjectName("SplashFrame")
        
        splash_layout = QVBoxLayout(self.frame)
        splash_layout.setContentsMargins(20, 20, 20, 20)
        
        title = QLabel("Axiom Automation Tool")
        title_font = QFont("Segoe UI", 16)
        title_font.setBold(True)
        title.setFont(title_font)
        title.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self.status_label = QLabel("Initializing...")
        self.status_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        
        self.progressBar = QProgressBar()
        self.progressBar.setTextVisible(False)

        splash_layout.addWidget(title)
        splash_layout.addStretch()
        splash_layout.addWidget(self.status_label)
        splash_layout.addWidget(self.progressBar)
        
        layout.addWidget(self.frame)

    def set_progress(self, value):
        self.progressBar.setValue(value)

# =============================================================================
# Browser Driver Class (NO CHANGES)
# =============================================================================
class BrowserDriver:
    def __init__(self, progress_callback):
        self.playwright = None
        self.browser = None
        self.page = None
        self.progress_callback = progress_callback
    def launch(self):
        self.progress_callback("Launching browser...")
        self.playwright = sync_playwright().start()
        self.browser = self.playwright.chromium.launch(headless=False, slow_mo=50)
        self.page = self.browser.new_page()
        self.progress_callback("Browser launched successfully.")
    def login(self, url, email, password):
        if not self.page: raise Exception("Browser is not launched.")
        self.progress_callback(f"Navigating to login page: {url}")
        self.page.goto(url, timeout=60000)
        email_locator = self.page.locator('#input28')
        expect(email_locator).to_be_visible(timeout=30000)
        email_locator.fill(email)
        next_button_locator = self.page.locator('input[value="Next"]')
        expect(next_button_locator).to_be_visible(timeout=30000)
        next_button_locator.click()
        password_locator = self.page.locator('#input29')
        expect(password_locator).to_be_visible(timeout=30000)
        password_locator.fill(password)
        verify_button_locator = self.page.locator('input[value="Verify"]')
        expect(verify_button_locator).to_be_visible(timeout=30000)
        verify_button_locator.click()
        try:
            push_option_locator = self.page.locator('[aria-label="Select to get a push notification to the Okta Verify app."]')
            expect(push_option_locator).to_be_visible(timeout=30000)
            push_option_locator.click()
            push_sent_locator = self.page.get_by_text("We've sent a push notification", exact=False)
            expect(push_sent_locator).to_be_visible(timeout=15000)
        except Exception: pass
        self.progress_callback(">>> Please approve the notification on your phone. <<<")
        self.page.wait_for_url(lambda url: "okta.com" not in url, timeout=120000)
        self.progress_callback("MFA approved. Login successful!")
        self.page.wait_for_timeout(3000)
    def wait_for_page_to_settle(self):
        self.page.wait_for_load_state("networkidle", timeout=30000)
    def navigate_to_create_user_page(self, base_url):
        create_url = f"{base_url}/CreateUserAccount"
        self.page.goto(create_url, timeout=60000)
        self.wait_for_page_to_settle()
    def fill_user_creation_form(self, user_details, user_password, postfix):
        username = f"{user_details['Username']}{postfix}"
        username_from_file = str(user_details['Username'])
        currency_code = username_from_file[:3]
        self.progress_callback(f"--- Creating user account: {username} ---")
        self.page.locator('//*[@id="accountForm"]/div/div[2]/div/div/input').click()
        self.page.locator("span", has_text="DEF (No Regulated Market)").click()
        self.page.locator('//*[@id="accountForm"]/div/div[3]/div/div/input').click()
        self.page.locator("span", has_text="Island Paradise Mobile (5007)").click()
        self.page.locator("#username").fill(username)
        self.page.locator("#password").fill(user_password)
        self.page.locator('//div[9]//div[1]//div[1]//input[1]').click()
        self.page.locator(f"span:has-text('({currency_code})')").click()
        self.page.locator("#submit").click()
        try:
            start_time = time.time()
            while time.time() - start_time < 30:
                if self.page.locator(".card-panel.green").is_visible():
                    self.progress_callback(f"  - Successfully created user: {username}")
                    return True
                if self.page.locator(".card-panel.red").is_visible():
                    self.page.locator("div.collapsible-header:has-text('Log')").click()
                    error_details = self.page.locator("#resultContainerMessage").inner_text()
                    self.progress_callback(f"  - Detailed Error: {error_details.strip()}")
                    return False
                self.page.wait_for_timeout(500)
            raise Exception("Timeout")
        except Exception as e:
            self.progress_callback(f"  - [CRITICAL] Could not determine creation status. Error: {e}")
            return False
    def migrate_user_to_lvc(self, base_url, username):
        user_accounts_url = f"{base_url}/useraccounts"
        self.page.goto(user_accounts_url, timeout=60000)
        self.wait_for_page_to_settle()
        self.page.locator("//button[@aria-label='Search']").click()
        self.page.locator("//input[@type='text']").first.fill(username)
        self.page.keyboard.press("Enter")
        self.page.wait_for_timeout(1500)
        self.wait_for_page_to_settle()
        user_row_locator = self.page.locator(f"tr:has-text('{username}')")
        expect(user_row_locator).to_be_visible(timeout=15000)
        three_dots_button = user_row_locator.locator("span.MuiIconButton-label")
        expect(three_dots_button).to_be_visible(timeout=10000)
        three_dots_button.click()
        migrate_option = self.page.get_by_role("menuitem", name="Migrate LVCS")
        expect(migrate_option).to_be_visible(timeout=10000)
        migrate_option.click()
        confirmation_button = self.page.get_by_role("button", name="Migrate")
        expect(confirmation_button).to_be_visible(timeout=10000)
        confirmation_button.click()
        self.page.wait_for_timeout(1000)
        self.progress_callback(f"Successfully migrated {username} to LVC.")
    def add_balance_to_user(self, base_url, username, amount):
        balance_url = f"{base_url}/BalanceUserAccount"
        self.page.goto(balance_url, timeout=60000)
        self.wait_for_page_to_settle()
        try:
            if self.page.locator("#loginName-error").is_visible(timeout=1000):
                self.progress_callback(f"  - [ERROR] User '{username}' not found. Skipping.")
                return
        except: pass
        self.page.locator("#loginName").fill(username)
        self.page.locator("#amount").fill(amount)
        self.page.locator("#submit").click()
        try:
            expect(self.page.locator(".card-panel.green")).to_be_visible(timeout=30000)
            self.progress_callback(f"  - Successfully added balance to {username}")
        except Exception as e:
            self.progress_callback(f"  - [CRITICAL] Could not determine balance status. Error: {e}")
    def close(self):
        if self.browser: self.browser.close()
        if self.playwright: self.playwright.stop()

# =============================================================================
# Helper Functions (NO CHANGES)
# =============================================================================
def parse_user_data(file_path, mode):
    lvc_users, standard_users = [], []
    file_extension = os.path.splitext(file_path)[1].lower()
    if file_extension in ['.xlsx', '.xls']:
        lvc_required_columns = ['lvc currency', 'username', 'lvc_username']
        standard_required_columns = ['std_currency', 'std_username']
        if mode in ["all", "lvc_only"]:
            lvc_df = pd.read_excel(file_path, sheet_name='Sheet1', header=0, usecols="A,B,D")
            lvc_df.columns = [str(col).strip().lower() for col in lvc_df.columns]
            if not all(col in lvc_df.columns for col in lvc_required_columns): raise ValueError("Missing LVC headers")
            lvc_df.rename(columns={'lvc currency': 'Currency', 'username': 'Username', 'lvc_username': 'LVC_username'}, inplace=True)
            lvc_df.dropna(subset=['Currency'], inplace=True)
            lvc_users = lvc_df.to_dict('records')
        if mode in ["all", "standard_only"]:
            standard_df = pd.read_excel(file_path, sheet_name='Sheet1', header=0, usecols="E,F")
            standard_df.columns = [str(col).strip().lower() for col in standard_df.columns]
            if not all(col in standard_df.columns for col in standard_required_columns): raise ValueError("Missing Standard headers")
            standard_df.rename(columns={'std_currency': 'Currency', 'std_username': 'Username'}, inplace=True)
            standard_df.dropna(subset=['Currency'], inplace=True)
            standard_users = standard_df.to_dict('records')
    elif file_extension == '.json':
        with open(file_path, 'r') as f: data = {k.lower(): v for k, v in json.load(f).items()}
        lvc_list_key = 'lvc_users' if 'lvc_users' in data else 'lvc_currencies'
        std_list_key = 'standard_users' if 'standard_users' in data else 'std_currencies'
        if mode in ["all", "lvc_only"] and lvc_list_key in data:
            for user in data[lvc_list_key]:
                norm_user = {k.lower().replace('_', ' '): v for k, v in user.items()}
                lvc_users.append({'Currency': norm_user.get('lvc currency', ''), 'Username': norm_user.get('username', ''), 'LVC_username': norm_user.get('lvc username', '')})
        if mode in ["all", "standard_only"] and std_list_key in data:
            for user in data[std_list_key]:
                norm_user = {k.lower().replace('_', ' '): v for k, v in user.items()}
                standard_users.append({'Currency': norm_user.get('std currency', ''), 'Username': norm_user.get('std username', '')})
    else: raise ValueError(f"Unsupported file type: {file_extension}")
    return lvc_users, standard_users
def parse_credentials_file(file_path):
    with open(file_path, 'r') as f: lines = f.readlines()
    if len(lines) < 2: raise ValueError("Credentials file needs 2 lines.")
    return lines[0].strip(), lines[1].strip()
def parse_gtp_list_file(file_path):
    with open(file_path, 'r') as f: data = json.load(f)
    if not isinstance(data, dict) or not data: raise ValueError("GTP list must be a non-empty JSON object.")
    return data
def update_excel_with_lvc_names(file_path, migrated_user, postfix):
    try:
        workbook = load_workbook(filename=file_path)
        sheet = workbook['Sheet1']
        original_username, original_currency = migrated_user['Username'], migrated_user['Currency']
        lvc_username = f"LVC_{original_username}{postfix}"
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row=row, column=1).value == original_currency and sheet.cell(row=row, column=2).value == original_username:
                sheet.cell(row=row, column=4).value = lvc_username
                break
        workbook.save(filename=file_path)
    except Exception as e: print(f"Error updating Excel: {e}")
def update_json_with_lvc_names(file_path, migrated_user, postfix):
    try:
        with open(file_path, 'r') as f: data = json.load(f)
        original_username, original_currency = migrated_user['Username'], migrated_user['Currency']
        lvc_username = f"LVC_{original_username}{postfix}"
        lvc_list_key = None
        for key in data:
            if key.lower() in ['lvc_users', 'lvc_currencies']: lvc_list_key = key; break
        if lvc_list_key:
            for user in data[lvc_list_key]:
                user_keys_lower = {k.lower().replace('_', ' '): k for k in user}
                currency_key, username_key = user_keys_lower.get('lvc currency'), user_keys_lower.get('username')
                if currency_key and username_key and user[currency_key] == original_currency and user[username_key] == original_username:
                    lvc_username_key = user_keys_lower.get('lvc username')
                    if lvc_username_key: user[lvc_username_key] = lvc_username; break
        with open(file_path, 'w') as f: json.dump(data, f, indent=2)
    except Exception as e: print(f"Error updating JSON: {e}")

# =============================================================================
# Automation Worker Class (NO CHANGES)
# =============================================================================
class AutomationWorker(QObject):
    progress_update = pyqtSignal(str)
    automation_error = pyqtSignal(str)
    automation_finished = pyqtSignal()
    def __init__(self, gtp_url, email, password, lvc_users, standard_users, user_password, user_data_path, mode, lvc_balance, standard_balance, postfix):
        super().__init__()
        self.gtp_url, self.email, self.password = gtp_url, email, password
        self.lvc_users, self.standard_users = lvc_users, standard_users
        self.user_password, self.user_data_path = user_password, user_data_path
        self.mode, self.lvc_balance, self.standard_balance = mode, lvc_balance, standard_balance
        self.postfix = postfix
        self.is_running = True
        self.driver = None
    def run(self):
        self.driver = BrowserDriver(progress_callback=self.progress_update.emit)
        try:
            self.driver.launch()
            if not self.is_running: return
            self.driver.login(self.gtp_url, self.email, self.password)
            if not self.is_running: return
            if self.mode in ["all", "lvc_only"]:
                created_lvc_users = []
                if self.lvc_users:
                    self.driver.navigate_to_create_user_page(self.gtp_url)
                    for user in self.lvc_users:
                        if not self.is_running: break
                        if self.driver.fill_user_creation_form(user, self.user_password, self.postfix):
                            created_lvc_users.append(user)
                if not self.is_running: return
                for user in created_lvc_users:
                    if not self.is_running: break
                    initial_username = f"{user['Username']}{self.postfix}"
                    self.driver.migrate_user_to_lvc(self.gtp_url, initial_username)
                    if self.user_data_path.lower().endswith(('.xlsx', '.xls')):
                        update_excel_with_lvc_names(self.user_data_path, user, self.postfix)
                    elif self.user_data_path.lower().endswith('.json'):
                        update_json_with_lvc_names(self.user_data_path, user, self.postfix)
                if not self.is_running: return
                for user in created_lvc_users:
                    if not self.is_running: break
                    lvc_username = f"LVC_{user['Username']}{self.postfix}"
                    self.driver.add_balance_to_user(self.gtp_url, lvc_username, self.lvc_balance)
            if self.mode in ["all", "standard_only"]:
                created_standard_users = []
                if self.standard_users:
                    self.driver.navigate_to_create_user_page(self.gtp_url)
                    for user in self.standard_users:
                        if not self.is_running: break
                        if self.driver.fill_user_creation_form(user, self.user_password, self.postfix):
                            created_standard_users.append(user)
                if not self.is_running: return
                for user in created_standard_users:
                    if not self.is_running: break
                    standard_username = f"{user['Username']}{self.postfix}"
                    self.driver.add_balance_to_user(self.gtp_url, standard_username, self.standard_balance)
            if self.is_running: self.progress_update.emit("\nAutomation complete!")
        except Exception as e:
            error_details = traceback.format_exc()
            full_error_message = f"A critical error stopped the automation.\n\nError Type: {type(e).__name__}\nDetails: {e}\n\nTraceback:\n{error_details}"
            self.progress_update.emit(f"[ERROR] {full_error_message}")
            self.automation_error.emit(full_error_message)
        finally:
            if self.driver: self.driver.close()
            self.automation_finished.emit()
    def stop(self): self.is_running = False

# =============================================================================
# Dialogs
# =============================================================================
class SettingsDialog(QDialog):
    def __init__(self, current_theme, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Settings")
        self.setFixedSize(300, 200)
        self.layout = QVBoxLayout()
        groupbox = QGroupBox("Theme")
        theme_layout = QVBoxLayout()
        self.radio_system = QRadioButton("Sync with System")
        self.radio_light = QRadioButton("Light Theme")
        self.radio_dark = QRadioButton("Dark Theme")
        if current_theme == "system": self.radio_system.setChecked(True)
        elif current_theme == "light": self.radio_light.setChecked(True)
        else: self.radio_dark.setChecked(True)
        theme_layout.addWidget(self.radio_system)
        theme_layout.addWidget(self.radio_light)
        theme_layout.addWidget(self.radio_dark)
        groupbox.setLayout(theme_layout)
        self.layout.addWidget(groupbox)
        self.ok_button = QPushButton("OK")
        self.ok_button.clicked.connect(self.accept)
        self.layout.addWidget(self.ok_button)
        self.setLayout(self.layout)
    def get_selected_theme(self):
        if self.radio_light.isChecked(): return "light"
        if self.radio_dark.isChecked(): return "dark"
        return "system"

class AboutDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("About")
        self.setFixedSize(400, 150)
        self.layout = QVBoxLayout()
        title = QLabel("Axiom Automation Tool")
        title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        title.setFont(QFont("Segoe UI", 12, QFont.Weight.Bold))
        # contact = QLabel("Contact: h.tavadarkar@zensar.com")
        contact = QLabel("Contact: "'<a href="mailto:h.tavadarkar">h.tavadarkar</a>')
        contact.setAlignment(Qt.AlignmentFlag.AlignCenter)
        contact.setOpenExternalLinks(True)
        version = QLabel("Version 1.4")
        version.setAlignment(Qt.AlignmentFlag.AlignCenter)
        link = QLabel('<a href="https://github.com/hrxtreme/AutomaticCurrencyCreator">Visit our GitHub Page</a>')
        link.setAlignment(Qt.AlignmentFlag.AlignCenter)
        link.setOpenExternalLinks(True)
        self.layout.addWidget(title)
        self.layout.addWidget(version)
        self.layout.addWidget(contact)
        self.layout.addWidget(link)
        self.setLayout(self.layout)

# =============================================================================
# Custom Title Bar
# =============================================================================
class CustomTitleBar(QWidget):
    def __init__(self, parent):
        super().__init__(parent)
        self.setObjectName("CustomTitleBar")
        self.parent = parent
        self.layout = QHBoxLayout()
        self.layout.setContentsMargins(10, 0, 0, 0)
        self.layout.setSpacing(0)
        
        # Title
        self.title_label = QLabel("Axiom Automation Tool")
        self.title_label.setFont(QFont("Segoe UI", 10, QFont.Weight.Bold))
        
        self.layout.addWidget(self.title_label)
        self.layout.addStretch()

        # Buttons
        self.settings_button = self.create_button("⚙️", "Settings")
        self.about_button = self.create_button("ℹ️", "About")
        self.minimize_button = self.create_button("—", "Minimize")
        self.maximize_button = self.create_button("🗖", "Maximize")
        self.close_button = self.create_button("✕", "Close")
        self.close_button.setObjectName("CloseButton")

        self.layout.addWidget(self.settings_button)
        self.layout.addWidget(self.about_button)
        self.layout.addWidget(self.minimize_button)
        self.layout.addWidget(self.maximize_button)
        self.layout.addWidget(self.close_button)

        self.setLayout(self.layout)
        self.pressing = False

    def create_button(self, text, tooltip):
        button = QPushButton(text)
        button.setFixedSize(45, 30)
        button.setToolTip(tooltip)
        button.setObjectName("TitleBarButton")
        return button
        
    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.start_pos = event.globalPosition().toPoint()
            self.pressing = True
    
    def mouseMoveEvent(self, event):
        if self.pressing and self.parent.isMaximized() == False:
            delta = event.globalPosition().toPoint() - self.start_pos
            self.parent.move(self.parent.pos() + delta)
            self.start_pos = event.globalPosition().toPoint()
            
    def mouseReleaseEvent(self, event):
        self.pressing = False

    def mouseDoubleClickEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.parent.toggle_maximize()

# =============================================================================
# Main Application Window (UI)
# =============================================================================
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setObjectName("MainWindow")
        self.setWindowFlags(Qt.WindowType.FramelessWindowHint)
        self.setMinimumSize(850, 800)

        self.automation_thread = None
        self.worker = None
        self.config = {}
        self.GTP_VERSIONS = {}

        self.central_widget = QWidget()
        self.setCentralWidget(self.central_widget)
        self.main_layout = QVBoxLayout(self.central_widget)
        self.main_layout.setContentsMargins(1, 1, 1, 1) 
        self.main_layout.setSpacing(0)

        self.title_bar = CustomTitleBar(self)
        self.main_layout.addWidget(self.title_bar)
        
        content_widget = QWidget()
        self.content_layout = QVBoxLayout(content_widget)
        self.content_layout.setSpacing(15)
        self.content_layout.setContentsMargins(20, 10, 20, 20)
        self.main_layout.addWidget(content_widget)
        
        self.load_config()
        self.setup_ui_elements()
        self.setup_connections()
        self.apply_theme()
        self.apply_config()
        self.check_start_button_state()
        self.is_maximized = False
        
    def setup_ui_elements(self):
        # 1. GTP Config
        config_card = QGroupBox("1. GTP Configuration")
        config_layout = QVBoxLayout()
        gtp_file_layout = QHBoxLayout()
        self.select_gtp_list_button = QPushButton("Select GTP List File (.json)")
        self.gtp_list_path_label = QLabel("No file selected.")
        gtp_file_layout.addWidget(self.select_gtp_list_button)
        gtp_file_layout.addWidget(self.gtp_list_path_label, 1)
        config_layout.addLayout(gtp_file_layout)
        self.gtp_dropdown = QComboBox()
        config_layout.addWidget(self.gtp_dropdown)
        config_card.setLayout(config_layout)
        self.content_layout.addWidget(config_card)

        # 2. Files & Settings
        files_card = QGroupBox("2. Files & Settings")
        files_grid_layout = QGridLayout()
        files_grid_layout.setSpacing(10)
        files_grid_layout.setColumnStretch(1, 1)
        files_grid_layout.setColumnStretch(3, 1)

        self.select_cred_button = QPushButton("Select Credentials File (.txt)")
        self.cred_path_label = QLabel("No file selected.")
        self.select_user_file_button = QPushButton("Select User Data File (excel or .json)")
        self.user_file_path_label = QLabel("No file selected.")
        files_grid_layout.addWidget(self.select_cred_button, 0, 0, 1, 2)
        files_grid_layout.addWidget(self.cred_path_label, 0, 2, 1, 2)
        files_grid_layout.addWidget(self.select_user_file_button, 1, 0, 1, 2)
        files_grid_layout.addWidget(self.user_file_path_label, 1, 2, 1, 2)
        
        user_pass_label = QLabel("New User Password:")
        self.user_password_input = QLineEdit()
        postfix_label = QLabel("Username Postfix:")
        self.postfix_input = QLineEdit()
        files_grid_layout.addWidget(user_pass_label, 2, 0)
        files_grid_layout.addWidget(self.user_password_input, 2, 1)
        files_grid_layout.addWidget(postfix_label, 3, 0)
        files_grid_layout.addWidget(self.postfix_input, 3, 1)

        lvc_balance_label = QLabel("LVC User Balance:")
        self.lvc_balance_input = QLineEdit()
        self.lvc_balance_input.setMaxLength(13)
        standard_balance_label = QLabel("Standard User Balance:")
        self.standard_balance_input = QLineEdit()
        self.standard_balance_input.setMaxLength(7)
        self.standard_balance_input.setValidator(QIntValidator(0, 9999999))
        files_grid_layout.addWidget(standard_balance_label, 2, 2)
        files_grid_layout.addWidget(self.standard_balance_input, 2, 3)
        files_grid_layout.addWidget(lvc_balance_label, 3, 2)
        files_grid_layout.addWidget(self.lvc_balance_input, 3, 3)

        files_card.setLayout(files_grid_layout)
        self.content_layout.addWidget(files_card)

        # 3. Processing Mode
        mode_card = QGroupBox("3. Processing Mode")
        mode_layout = QHBoxLayout()
        self.radio_all = QRadioButton("LVC + Standard")
        self.radio_lvc = QRadioButton("Only LVC")
        self.radio_standard = QRadioButton("Only Standard")
        self.radio_all.setChecked(True)
        mode_layout.addStretch(1)
        mode_layout.addWidget(self.radio_all)
        mode_layout.addStretch(1)
        mode_layout.addWidget(self.radio_lvc)
        mode_layout.addStretch(1)
        mode_layout.addWidget(self.radio_standard)
        mode_layout.addStretch(1)
        mode_card.setLayout(mode_layout)
        self.content_layout.addWidget(mode_card)

        # Start Button
        self.start_button = QPushButton("▶ Start Automation")
        self.start_button.setObjectName("StartButton")
        self.start_button.setMinimumHeight(40)
        self.content_layout.addWidget(self.start_button)
        
        # Status Log
        log_card = QGroupBox("Status Log")
        log_layout = QVBoxLayout()
        self.status_log = QTextEdit()
        self.status_log.setReadOnly(True)
        log_layout.addWidget(self.status_log)
        log_card.setLayout(log_layout)
        self.content_layout.addWidget(log_card)
        self.content_layout.setStretch(4, 1) # Make log expand

    def setup_connections(self):
        # Title Bar
        self.title_bar.minimize_button.clicked.connect(self.showMinimized)
        self.title_bar.maximize_button.clicked.connect(self.toggle_maximize)
        self.title_bar.close_button.clicked.connect(self.close)
        self.title_bar.settings_button.clicked.connect(self.open_settings_dialog)
        self.title_bar.about_button.clicked.connect(self.open_about_dialog)

        # Main Content
        self.select_gtp_list_button.clicked.connect(self.select_gtp_list_file)
        self.select_cred_button.clicked.connect(self.select_credentials_file)
        self.select_user_file_button.clicked.connect(self.select_user_data_file)
        self.start_button.clicked.connect(self.start_automation)

    def start_automation(self):
        gtp_selection = self.gtp_dropdown.currentText()
        if not gtp_selection:
             QMessageBox.warning(self, "Input Error", "Please select a GTP version.")
             return
        gtp_url = self.GTP_VERSIONS[gtp_selection]
        cred_path, user_data_path = self.cred_path_label.text(), self.user_file_path_label.text()
        user_password, postfix = self.user_password_input.text(), self.postfix_input.text()
        lvc_balance, standard_balance = self.lvc_balance_input.text(), self.standard_balance_input.text()
        
        mode = "all"
        if self.radio_lvc.isChecked(): mode = "lvc_only"
        elif self.radio_standard.isChecked(): mode = "standard_only"

        errors = []
        if "No file selected" in cred_path: errors.append("Credentials file not selected.")
        if "No file selected" in user_data_path: errors.append("User data file not selected.")
        if not user_password: errors.append("New user password cannot be empty.")
        if not postfix: errors.append("Postfix cannot be empty.")
        if not lvc_balance.isdigit() or not standard_balance.isdigit(): errors.append("Balance amounts must be valid numbers.")
        if errors:
            QMessageBox.warning(self, "Input Error", "\n".join(errors)); return

        self.log_message("="*50 + f"\nStarting new run in '{mode}' mode...")
        try:
            email, password = parse_credentials_file(cred_path)
            lvc_users, standard_users = parse_user_data(user_data_path, mode)
            self.log_message(f"  - Validation successful: Found {len(lvc_users)} LVC and {len(standard_users)} Standard users.")
        except Exception as e:
            error_message = f"Failed to read input file.\n\nError: {e}"
            self.log_message(f"[ERROR] {error_message}")
            QMessageBox.critical(self, "File Error", error_message); return

        self.toggle_controls(False)
        self.automation_thread = QThread()
        self.worker = AutomationWorker(gtp_url, email, password, lvc_users, standard_users, user_password, user_data_path, mode, lvc_balance, standard_balance, postfix)
        self.worker.moveToThread(self.automation_thread)
        self.worker.progress_update.connect(self.log_message)
        self.worker.automation_error.connect(self.on_automation_error)
        self.worker.automation_finished.connect(self.on_automation_finished)
        self.automation_thread.started.connect(self.worker.run)
        self.automation_thread.finished.connect(self.automation_thread.deleteLater)
        self.automation_thread.start()

    def on_automation_error(self, error_message):
        QMessageBox.critical(self, "Automation Error", error_message)
    def on_automation_finished(self):
        self.log_message("Process finished.")
        if self.automation_thread:
            self.automation_thread.quit()
            self.automation_thread.wait()
        self.toggle_controls(True)
        self.automation_thread, self.worker = None, None
    def select_gtp_list_file(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "Select GTP List File", "", "JSON Files (*.json)")
        if file_path: self.load_gtp_list_from_path(file_path, save=True)
    def select_credentials_file(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "Select Credentials File", "", "Text Files (*.txt)")
        if file_path:
            self.cred_path_label.setText(file_path)
            self.config['credentials_path'] = file_path
            self.save_config(); self.check_start_button_state()
    def select_user_data_file(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "Select User Data File", "", "Data Files (*.xlsx *.xls *.json)")
        if file_path:
            self.user_file_path_label.setText(file_path)
            self.config['user_data_path'] = file_path
            self.save_config(); self.check_start_button_state()
    def load_gtp_list_from_path(self, file_path, save=False):
        try:
            self.GTP_VERSIONS = parse_gtp_list_file(file_path)
            self.gtp_dropdown.clear(); self.gtp_dropdown.addItems(self.GTP_VERSIONS.keys())
            self.gtp_list_path_label.setText(file_path)
            if save: self.config['gtp_list_path'] = file_path; self.save_config()
        except Exception as e:
            self.GTP_VERSIONS = {}; self.gtp_dropdown.clear()
            QMessageBox.critical(self, "File Error", f"Failed to load GTP list file: {e}")
        self.check_start_button_state()
    def load_config(self):
        if os.path.exists(CONFIG_FILE):
            with open(CONFIG_FILE, 'r') as f: self.config = json.load(f)
        else: self.config = {}
    def save_config(self):
        self.config['postfix'] = self.postfix_input.text()
        self.config['user_password'] = self.user_password_input.text()
        self.config['last_gtp_selection'] = self.gtp_dropdown.currentText()
        self.config['lvc_balance'] = self.lvc_balance_input.text()
        self.config['standard_balance'] = self.standard_balance_input.text()
        with open(CONFIG_FILE, 'w') as f: json.dump(self.config, f, indent=4)
    def apply_config(self):
        gtp_path = self.config.get('gtp_list_path')
        if gtp_path and os.path.exists(gtp_path):
            self.load_gtp_list_from_path(gtp_path, save=False)
            last_selection = self.config.get('last_gtp_selection')
            if last_selection:
                index = self.gtp_dropdown.findText(last_selection)
                if index != -1: self.gtp_dropdown.setCurrentIndex(index)
        cred_path = self.config.get('credentials_path')
        if cred_path and os.path.exists(cred_path): self.cred_path_label.setText(cred_path)
        user_data_path = self.config.get('user_data_path')
        if user_data_path and os.path.exists(user_data_path): self.user_file_path_label.setText(user_data_path)
        self.postfix_input.setText(self.config.get('postfix', 'x1'))
        self.user_password_input.setText(self.config.get('user_password', 'snow'))
        self.lvc_balance_input.setText(self.config.get('lvc_balance', '7000000000000'))
        self.standard_balance_input.setText(self.config.get('standard_balance', '9999999'))
    def check_start_button_state(self):
        gtp_loaded = bool(self.GTP_VERSIONS)
        creds_loaded = "No file selected" not in self.cred_path_label.text()
        users_loaded = "No file selected" not in self.user_file_path_label.text()
        self.start_button.setEnabled(gtp_loaded and creds_loaded and users_loaded)
    def log_message(self, message):
        self.status_log.append(message); logging.info(message)
    def toggle_controls(self, enabled):
        for widget in self.central_widget.findChildren(QWidget):
            if isinstance(widget, (QPushButton, QLineEdit, QComboBox, QRadioButton)) and widget not in self.title_bar.findChildren(QWidget):
                widget.setEnabled(enabled)
    def toggle_maximize(self):
        if self.is_maximized:
            self.showNormal()
            self.title_bar.maximize_button.setText("🗖")
            self.is_maximized = False
        else:
            self.showMaximized()
            self.title_bar.maximize_button.setText("🗗")
            self.is_maximized = True
    def open_settings_dialog(self):
        current_theme = self.config.get("theme", "system")
        dialog = SettingsDialog(current_theme, self)
        if dialog.exec():
            new_theme = dialog.get_selected_theme()
            if new_theme != current_theme:
                self.config["theme"] = new_theme
                self.save_config()
                self.apply_theme()
    def open_about_dialog(self):
        dialog = AboutDialog(self)
        dialog.exec()
    def apply_theme(self):
        theme = self.config.get("theme", "system")
        stylesheet = ""
        if theme == "system":
            stylesheet = DARK_THEME if is_windows_dark_theme() else LIGHT_THEME
        elif theme == "dark":
            stylesheet = DARK_THEME
        else: # light
            stylesheet = LIGHT_THEME
        self.setStyleSheet(stylesheet)
    def closeEvent(self, event):
        self.save_config()
        if self.automation_thread and self.automation_thread.isRunning():
            self.worker.stop(); self.automation_thread.quit(); self.automation_thread.wait()
        event.accept()

def main():
    app = QApplication(sys.argv)
    
    splash = SplashScreen()
    config = {}
    if os.path.exists(CONFIG_FILE):
        with open(CONFIG_FILE, 'r') as f: config = json.load(f)
    theme = config.get("theme", "system")
    app.setStyleSheet(DARK_THEME if (theme == "dark" or (theme == "system" and is_windows_dark_theme())) else LIGHT_THEME)
    splash.show()
    
    main_window = None
    # Apply theme before showing window
    QApplication.processEvents()
    main_window = MainWindow()
    main_window.apply_theme() # Apply theme after construction
    
    for i in range(101):
        splash.set_progress(i)
        time.sleep(0.01)
        QApplication.processEvents()

    splash.close()
    if main_window: main_window.show()
    sys.exit(app.exec())

if __name__ == '__main__':
    main()

